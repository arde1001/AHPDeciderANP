"""Interchange export: the whole model as plain tables (Excel workbook or CSV zip).

Meant to be read by people and by other AHP/ANP software (SuperDecisions,
Expert Choice / Comparion, SpiceLogic, R, Python...), none of which share a
file format. So everything is exported in the most common shapes:

* ``Structure``: hierarchy (AHP) or clusters/nodes (ANP) with ids.
* ``Connections`` (ANP): 0/1 node matrix, row i / column j = 1 when i is
  compared with respect to j (in SuperDecisions terms: node j connects to i).
* ``Judgments``: one row per pair: A, B, numeric a_AB, preferred side and
  intensity 1-9, i.e. what you type into a questionnaire / verbal mode.
* one full reciprocal matrix per comparison (labels in the first row and
  column, priorities in the last column), the form matrix-entry modes and
  scripts expect.
* ``Comparisons``: index of the matrices with lambda_max, CI, RI, CR.
* ``Results`` and, for ANP, the four supermatrices, to cross-check other tools.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import numpy as np

from .. import __version__
from .ahp import GOAL_ID, AHPModel
from .anp import ANPModel
from .pairwise import (
    CR_THRESHOLD,
    ComparisonSet,
    analyze,
    describe_judgment,
    format_judgment,
)

APP_NAME = "AHP / ANP Decision Maker"


@dataclass
class MatrixTable:
    code: str  # M01, M02, ...
    title: str
    with_respect_to: str
    compared: str  # what is compared: "Criteria", "Alternatives", a cluster name, "Clusters"
    short: str  # short label for sheet / file names
    group: str
    names: list[str]
    matrix: np.ndarray
    priorities: np.ndarray
    lambda_max: float
    ci: float
    ri: float
    cr: float | None
    judged: int
    pairs: int


@dataclass
class Grid:
    """A labelled square matrix (connections, supermatrices)."""

    name: str
    labels: list[str]
    values: np.ndarray
    note: str = ""


@dataclass
class ExportTables:
    about: list[list]
    structure: list[list]
    judgments: list[list]
    index: list[list]
    matrices: list[MatrixTable]
    results: list[list]
    grids: list[Grid] = field(default_factory=list)  # ANP: connections + supermatrices
    scores: list[list] = field(default_factory=list)  # AHP: direct values and ratings, if any


# ---- building the tables ----------------------------------------------------


def _describe(model, s: ComparisonSet) -> tuple[str, str, str]:
    """(with respect to, what is compared, short label) of a comparison set."""
    if isinstance(model, AHPModel):
        if s.key.startswith("levels:"):
            name = model.find(s.key[len("levels:"):]).name
            return name, "Rating levels", f"{name} - rating levels"
        node = model.find(s.key)
        compared = "Alternatives" if not node.children else ("Criteria" if node.id == GOAL_ID else "Sub-criteria")
        return node.name, compared, f"{node.name} - {compared}"
    kind, *ids = s.key.split(":")
    if kind == "cluster":
        name = model.cluster(ids[0]).name
        return f"cluster {name}", "Clusters", f"Clusters wrt {name}"
    node, cluster = model.node(ids[0]), model.cluster(ids[1])
    return f"{model.cluster_of(ids[0]).name}: {node.name}", cluster.name, f"{node.name} - {cluster.name}"


def _matrices(model) -> list[MatrixTable]:
    out = []
    for k, s in enumerate(model.comparison_sets(), 1):
        comp = model.comparison(s.key)
        m = comp.matrix(s.item_ids)
        r = analyze(m)
        n = len(s.item_ids)
        wrt, compared, short = _describe(model, s)
        out.append(MatrixTable(
            code=f"M{k:02d}", title=s.title, with_respect_to=wrt, compared=compared, short=short, group=s.group,
            names=list(s.item_names), matrix=m, priorities=r.priorities, lambda_max=r.lambda_max,
            ci=r.ci, ri=r.ri, cr=r.cr, judged=comp.count_set(s.item_ids), pairs=n * (n - 1) // 2,
        ))
    return out


def _about(model, method: str) -> list[list]:
    tips = [
        ["Import tips", ""],
        ["Expert Choice / Comparion, SpiceLogic, questionnaires",
         "Recreate the hierarchy from 'Structure', then enter each pair from 'Judgments' "
         "(Preferred + Intensity), or type the matrices in numerical/matrix mode."],
        ["SuperDecisions",
         "Create the clusters and nodes from 'Structure' and the links from 'Connections' "
         "(row i, column j = 1: node j connects to node i). Enter each node comparison and the cluster "
         "comparisons from the matrix sheets/files; compare the supermatrices with SuperDecisions' own."],
        ["Direct values and ratings",
         "Criteria scored without pairwise comparison are listed in 'Scores' (measured value or rating per "
         "alternative, and the resulting priority). Enter them in the other tool's data/ratings mode; the rating "
         "levels' own pairwise matrix is among the matrix sheets."],
        ["Spreadsheets, R, Python",
         "Every matrix has item names in the first row and column; read it with the first column as "
         "row labels (e.g. pandas.read_csv(f, index_col=0)). Drop the 'Priority' column to get the "
         "square matrix."],
    ]
    return [
        ["Exported by", f"{APP_NAME} {__version__}"],
        ["Export date", date.today().isoformat()],
        ["Decision", model.goal],
        ["Method", method],
        [],
        ["Conventions", ""],
        ["Scale", "Saaty 1-9 with reciprocals: 1 equal, 3 moderate, 5 strong, 7 very strong, 9 extreme; "
                  "2, 4, 6, 8 in between; 1/x when the second item dominates."],
        ["Matrices", "Cell (row i, column j) = how strongly row item i dominates column item j. "
                     "a_ji = 1 / a_ij, diagonal = 1."],
        ["Unjudged pairs", "Count as 1 (equal); see the 'Judged' column."],
        ["Priorities", "Normalized principal right eigenvector of each matrix."],
        ["Consistency", f"CI = (lambda_max - n)/(n - 1), CR = CI/RI (Saaty's RI); CR <= {CR_THRESHOLD:.2f} "
                        "is acceptable. Not defined for n < 3."],
        [],
        *tips,
    ]


def _judgments(matrices: list[MatrixTable], model) -> list[list]:
    rows = [[
        "Matrix", "Comparison", "With respect to", "Compared", "Item A", "Item B", "Value (A over B)", "Value as text",
        "Preferred", "Intensity (1-9)", "Verbal", "Judged",
    ]]
    sets = {s.title: s for s in model.comparison_sets()}
    for t in matrices:
        s = sets[t.title]
        comp = model.comparison(s.key)
        n = len(t.names)
        for i in range(n):
            for j in range(i + 1, n):
                v = float(t.matrix[i, j])
                preferred = "Equal" if abs(v - 1) < 1e-9 else (t.names[i] if v > 1 else t.names[j])
                rows.append([
                    t.code, t.title, t.with_respect_to, t.compared, t.names[i], t.names[j], v, format_judgment(v),
                    preferred, round(max(v, 1 / v)), describe_judgment(v),
                    "yes" if comp.is_set(s.item_ids[i], s.item_ids[j]) else "no (default 1)",
                ])
    return rows


def _index(matrices: list[MatrixTable]) -> list[list]:
    rows = [["Matrix", "Comparison", "With respect to", "Compared", "Group", "Items", "Judged", "lambda_max", "CI", "RI",
             "CR", "Consistent"]]
    for t in matrices:
        rows.append([
            t.code, t.title, t.with_respect_to, t.compared, t.group, len(t.names), f"{t.judged}/{t.pairs}",
            t.lambda_max, t.ci, t.ri, "" if t.cr is None else t.cr,
            "yes" if t.cr is None or t.cr <= CR_THRESHOLD else "no",
        ])
    return rows


def _ahp_tables(model: AHPModel) -> tuple[list[list], list[list], list[Grid]]:
    structure = [["ID", "Type", "Name", "Parent ID", "Parent", "Level", "Path", "Alternatives scored by"]]
    for node, parent, depth in model.walk():
        kind = "Goal" if parent is None else ("Criterion" if depth == 1 else "Sub-criterion")
        structure.append([
            node.id, kind, node.name, parent.id if parent else "", parent.name if parent else "", depth,
            " > ".join(model.path(node.id)), "" if node.children else _SCORING_LABEL[model.scoring_mode(node)],
        ])
    for a in model.alternatives:
        structure.append([a.id, "Alternative", a.name, "", "", "", "", ""])

    dist = model.evaluate(synthesis="distributive")
    ideal = model.evaluate(synthesis="ideal")
    chosen = dist if model.synthesis == "distributive" else ideal
    results = [["Alternatives"], ["Rank", "Alternative", "Priority (distributive)", "Priority (ideal)"]]
    for rank, k in enumerate(chosen.ranking(), 1):
        results.append([rank, dist.alternative_names[k], float(dist.alternative_scores[k]),
                        float(ideal.alternative_scores[k])])
    results += [[], [f"Ranked by the project's synthesis mode: {model.synthesis}"], [],
                ["Criteria"], ["Path", "Level", "Local weight", "Global weight", "Leaf"]]
    for node, _parent, depth in model.walk():
        if node.id == GOAL_ID:
            continue
        results.append([" > ".join(model.path(node.id)), depth, dist.local_weights[node.id],
                        dist.global_weights[node.id], "yes" if not node.children else ""])
    return structure, results, []


_SCORING_LABEL = {"pairwise": "pairwise comparison", "direct": "direct values", "ratings": "ratings"}


def _scores(model: AHPModel) -> list[list]:
    """Measured values and ratings of the leaves that don't use pairwise comparison."""
    result = model.evaluate()
    rows = [["Criterion", "Scoring", "Direction", "Unit", "Alternative", "Value", "Rating", "Rating weight (ideal)",
             "Priority under criterion"]]
    for leaf in model.leaves():
        mode = model.scoring_mode(leaf)
        if mode == "pairwise":
            continue
        inp = model.leaf_inputs[leaf.id]
        path = " > ".join(model.path(leaf.id)) or model.goal
        levels = model.level_priorities(leaf.id) if mode == "ratings" else None
        ideal = {}
        if levels is not None:
            top = levels.priorities.max()
            ideal = {lv.id: float(p / top) for lv, p in zip(inp.levels, levels.priorities, strict=True)}
        names = {lv.id: lv.name for lv in inp.levels}
        for k, a in enumerate(model.alternatives):
            level = inp.ratings.get(a.id)
            rows.append([
                path, _SCORING_LABEL[mode],
                ("higher is better" if inp.direction == "benefit" else "lower is better") if mode == "direct" else "",
                inp.unit if mode == "direct" else "",
                a.name,
                inp.values.get(a.id, "") if mode == "direct" else "",
                names.get(level, "") if mode == "ratings" else "",
                ideal.get(level, "") if mode == "ratings" else "",
                float(result.local[leaf.id].priorities[k]),
            ])
    return rows if len(rows) > 1 else []


def _anp_tables(model: ANPModel) -> tuple[list[list], list[list], list[Grid]]:
    alt = model.effective_alternatives_cluster()
    start = model.effective_start_node()
    structure = [["Cluster ID", "Cluster", "Node ID", "Node", "Alternatives cluster", "Start (goal) node"]]
    for n, c in model.nodes():
        structure.append([c.id, c.name, n.id, n.name, "yes" if alt and c.id == alt.id else "",
                          "yes" if n.id == start else ""])
    for c in model.clusters:
        if not c.nodes:
            structure.append([c.id, c.name, "", "", "", ""])

    r = model.evaluate()
    labels = [f"{c}: {n}" for c, n in zip(r.node_clusters, r.node_names, strict=True)]
    links = np.array([[1 if model.is_linked(j, i) else 0 for j in r.node_ids] for i in r.node_ids], dtype=int)
    grids = [
        Grid("Connections", labels, links,
             "1 in row i, column j: node i is compared with respect to node j (node j connects to node i)."),
        Grid("Unweighted supermatrix", labels, r.unweighted, "Columns: local priorities w.r.t. each node."),
        Grid("Cluster matrix", r.cluster_names, r.cluster_matrix, "Column D: cluster weights w.r.t. cluster D."),
        Grid("Weighted supermatrix", labels, r.weighted,
             "Blocks times cluster weights, columns normalized; sinks made absorbing (1 on the diagonal)."),
        Grid("Limit supermatrix", labels, r.limit, f"Limit of the weighted supermatrix: {r.limit_status}."),
    ]
    results = [["Alternatives"], ["Rank", "Alternative", "Priority"]]
    for rank, k in enumerate(r.ranking(), 1):
        results.append([rank, r.alternative_names[k], float(r.alternative_scores[k])])
    results += [[], [f"Read from: {model.node(start).name if start else 'average of all columns'}; "
                     f"limit matrix {r.limit_status}."], [],
                ["Limit priorities"], ["Cluster", "Node", "Limit priority", "Normalized by cluster"]]
    for k in range(len(r.node_ids)):
        norm = r.normalized[k]
        results.append([r.node_clusters[k], r.node_names[k], float(r.priorities[k]),
                        "" if np.isnan(norm) else float(norm)])
    return structure, results, grids


def build_tables(model: AHPModel | ANPModel) -> ExportTables:
    matrices = _matrices(model)
    if isinstance(model, AHPModel):
        structure, results, grids = _ahp_tables(model)
        scores = _scores(model)
        method = "AHP (Analytic Hierarchy Process)"
    else:
        structure, results, grids = _anp_tables(model)
        scores = []
        method = "ANP (Analytic Network Process)"
    return ExportTables(
        about=_about(model, method), structure=structure, judgments=_judgments(matrices, model),
        index=_index(matrices), matrices=matrices, results=results, grids=grids, scores=scores,
    )


def matrix_rows(t: MatrixTable) -> list[list]:
    """Header row of item names, then one row per item; last column = priority."""
    rows = [["Row over column", *t.names, "Priority"]]
    for i, name in enumerate(t.names):
        rows.append([name, *(float(x) for x in t.matrix[i]), float(t.priorities[i])])
    return rows


def grid_rows(g: Grid) -> list[list]:
    cast = int if np.issubdtype(g.values.dtype, np.integer) else float
    return [["", *g.labels], *([label, *(cast(x) for x in g.values[i])] for i, label in enumerate(g.labels))]


# ---- writers -----------------------------------------------------------------


def _slug(text: str, limit: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")[:limit] or "x"


def _csv_text(rows: list[list]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    for row in rows:
        writer.writerow([f"{v:.10g}" if isinstance(v, float) else v for v in row])
    return buf.getvalue()


def matrix_filename(t: MatrixTable) -> str:
    return f"matrices/{t.code}_{_slug(t.short)}.csv"


def export_csv_zip(model: AHPModel | ANPModel, path: str | Path) -> None:
    """A zip of UTF-8 (with BOM, for Excel) CSV files plus a README."""
    tables = build_tables(model)
    index = [tables.index[0] + ["File"]] + [
        row + [matrix_filename(t)] for row, t in zip(tables.index[1:], tables.matrices, strict=True)
    ]
    files = {
        "structure.csv": tables.structure,
        "judgments.csv": tables.judgments,
        "comparisons.csv": index,
        "results.csv": tables.results,
    }
    if tables.scores:
        files["scores.csv"] = tables.scores
    for t in tables.matrices:
        files[matrix_filename(t)] = matrix_rows(t)
    for g in tables.grids:
        files[f"{_slug(g.name.lower())}.csv"] = grid_rows(g)
    lines = []
    for row in tables.about:
        cells = [str(c) for c in row if c != ""]
        lines.append(": ".join(cells))
    readme = "\n".join(lines) + "\n\nFiles:\n" + "".join(f"  {name}\n" for name in sorted(files))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README.txt", readme)
        for name, rows in files.items():
            z.writestr(name, "﻿" + _csv_text(rows))


def _sheet_name(text: str, used: set[str]) -> str:
    base = re.sub(r"\s+", " ", re.sub(r"[\[\]:*?/\\]", " ", text)).strip()[:31].strip() or "Sheet"
    name, k = base, 2
    while name.lower() in used:
        suffix = f" ({k})"
        name, k = base[: 31 - len(suffix)] + suffix, k + 1
    used.add(name.lower())
    return name


def export_xlsx(model: AHPModel | ANPModel, path: str | Path) -> None:
    """One workbook: About, Structure, Judgments, Comparisons, one sheet per matrix, Results, supermatrices."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    tables = build_tables(model)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    wb = Workbook()
    used: set[str] = set()

    def sheet(title: str, rows: list[list], header: bool = True, widths: int = 60, reserved: bool = False):
        ws = wb.create_sheet(title if reserved else _sheet_name(title, used))
        for row in rows:
            ws.append(row)
        if header and rows:
            for cell in ws[1]:
                cell.font = bold
                cell.fill = head_fill
            ws.freeze_panes = "A2"
        for col in ws.columns:
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=4)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(widths, max(8, longest + 2))
        return ws

    about = sheet("About", tables.about, header=False, widths=110)
    for row in about.iter_rows(min_col=1, max_col=1):
        row[0].font = bold
    sheet("Structure", tables.structure)
    if tables.grids:
        g = tables.grids[0]
        ws = sheet(g.name, grid_rows(g), widths=28)
        ws.freeze_panes = "B2"
        ws.append([])
        ws.append([g.note])
    sheet("Judgments", tables.judgments)
    # reserve the matrix sheet names first so the index can link to them
    sheet_names = [_sheet_name(f"{t.code} {t.short}", used) for t in tables.matrices]
    index = [tables.index[0] + ["Sheet"]]
    for row, name in zip(tables.index[1:], sheet_names, strict=True):
        index.append(row + [name])
    ws_index = sheet("Comparisons", index)
    for t, name in zip(tables.matrices, sheet_names, strict=True):
        rows = matrix_rows(t)
        ws = sheet(name, rows, widths=28, reserved=True)
        ws.freeze_panes = "B2"
        n = len(t.names)
        for r in range(2, n + 2):
            ws.cell(r, 1).font = bold
            for c in range(2, n + 2):
                ws.cell(r, c).number_format = "# ?/?"  # shows 1/3, 5, ... while staying numeric
            ws.cell(r, n + 2).number_format = "0.0000"
        ws.append([])
        for label, value in [
            ("Comparison", t.title), ("With respect to", t.with_respect_to), ("Compared", t.compared),
            ("Judged pairs", f"{t.judged}/{t.pairs}"), ("lambda_max", t.lambda_max), ("CI", t.ci), ("RI", t.ri),
            ("CR", "n < 3: always consistent" if t.cr is None else t.cr),
        ]:
            ws.append([label, value])
            ws.cell(ws.max_row, 1).font = bold
            if isinstance(value, float):
                ws.cell(ws.max_row, 2).number_format = "0.0000"
    for r in range(2, ws_index.max_row + 1):
        for c in range(1, ws_index.max_column + 1):
            if isinstance(ws_index.cell(r, c).value, float):
                ws_index.cell(r, c).number_format = "0.0000"
        target = ws_index.cell(r, ws_index.max_column)
        target.hyperlink = "#'{}'!A1".format(str(target.value).replace("'", "''"))
        target.style = "Hyperlink"
    if tables.scores:
        sheet("Scores", tables.scores)
    sheet("Results", tables.results, header=False)
    for g in tables.grids[1:]:
        ws = sheet(g.name, grid_rows(g), widths=28)
        ws.freeze_panes = "B2"
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0.0000"
        ws.append([])
        ws.append([g.note])
    wb.remove(wb.worksheets[0])  # the default empty sheet
    wb.save(path)
